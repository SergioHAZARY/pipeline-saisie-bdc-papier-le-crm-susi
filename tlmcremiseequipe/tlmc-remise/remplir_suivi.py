#!/usr/bin/env python3
"""Remplit la colonne F (Montant endossé) du « Suivi saisie paiement SUSI » à partir du rapport TLMC.

Entrées :
  - le rapport produit par `agent_ftp.py finaliser` (rapport.csv / rapport_sheet.csv) :
    colonnes « nom du fichier », « montant endossé (€) », « lien vers le fichier TLMC », « commentaire » ;
  - le classeur de suivi : export xlsx du Google Sheet (mode `xlsx`) OU le Google Sheet lui-même
    via gspread (mode `gsheet`, nécessite un client OAuth « Desktop » dans ~/.config/gspread/credentials.json).

Règles :
  - onglets « AA JJ-MM » d'abord (celui du jour du lot, puis les autres AA), puis « OUT JJ-MM » ;
  - colonne A = nom du fichier (rapprochement tolérant : .pdf, OUT/ATLAS/III/FR, CH02 = CH2, 6082026 = 06082026) ;
  - F = montant endossé ; si F ≠ E (montant saisi dans SUSI) → ligne en rouge ;
  - la colonne « Lien »/« LINK » (I) reçoit le lien TLMC si elle est vide ;
  - les lots introuvables sont listés (à ajouter à la main dans l'onglet OUT du jour).

Usage :
  .venv/bin/python remplir_suivi.py --rapport sessions/agent_ftp/rapport_sheet.csv \
      --xlsx suivi.xlsx --sortie suivi_rempli.xlsx [--json patch.json]
  .venv/bin/python remplir_suivi.py --rapport ... --gsheet 1exk9CYqd0lYv... [--dry-run]
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from dataclasses import dataclass, field, asdict
from pathlib import Path

ROUGE = "FF0000"
COL_NOM, COL_SUSI, COL_ENDOSSE, COL_COHERENCE, COL_LIEN = 1, 5, 6, 7, 9  # A, E, F, G, I (1-based)
MOTS_IGNORES = {"FR", "ATLAS", "III", "II", "OUT", "AA", "PDF"}


# ----------------------------------------------------------------------------- normalisation
def cle_lot(nom: str) -> tuple | None:
    """Clé de rapprochement d'un nom de lot : (date JJMMAAAA, catégorie, nb chèques, index)."""
    if not nom:
        return None
    s = str(nom).strip().upper()
    s = re.sub(r"\.PDF$", "", s)
    s = re.sub(r"[_\-]", " ", s)
    m = re.match(r"\s*(\d{7,8})\s*(.*)$", s)
    if not m:
        return None
    date = m.group(1).zfill(8)
    reste = m.group(2)
    mch = re.search(r"(\d{1,3})\s*CH\s*0*(\d*)", reste)
    if not mch:
        return None
    nb = int(mch.group(1))
    idx = int(mch.group(2)) if mch.group(2) else 0
    reste = (reste[: mch.start()] + " " + reste[mch.end():])
    mots = [w for w in re.split(r"[^A-Z]+", reste) if w and w not in MOTS_IGNORES]
    categorie = " ".join(sorted(mots))
    return (date, categorie, nb, idx)


def montant(v) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    s = str(v).strip()
    if s.startswith("="):
        return None  # formule : valeur inconnue hors Google Sheets
    s = s.replace("€", "").replace(" ", "").replace("\xa0", "").replace(" ", "")
    s = s.replace(",", ".")
    try:
        return round(float(s), 2)
    except ValueError:
        return None


def date_onglet(titre: str) -> str | None:
    m = re.match(r"\s*(AA|OUT)\s+(\d{2})-(\d{2})", titre.strip(), re.I)
    return (m.group(2) + m.group(3)) if m else None


def famille_onglet(titre: str) -> str | None:
    m = re.match(r"\s*(AA|OUT)\b", titre.strip(), re.I)
    return m.group(1).upper() if m else None


# ----------------------------------------------------------------------------- rapport
@dataclass
class Lot:
    fichier: str
    montant: float
    lien: str
    commentaire: str
    cle: tuple | None = None
    resultat: dict = field(default_factory=dict)


def lire_rapport(chemin: Path) -> list[Lot]:
    lots = []
    with chemin.open(encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            nom = row.get("nom du fichier") or row.get("fichier") or ""
            if not nom.strip():
                continue
            lot = Lot(
                fichier=nom.strip(),
                montant=montant(row.get("montant endossé (€)") or row.get("montant")) or 0.0,
                lien=(row.get("lien vers le fichier TLMC") or row.get("lien") or "").strip(),
                commentaire=(row.get("commentaire") or "").strip(),
            )
            lot.cle = cle_lot(lot.fichier)
            lots.append(lot)
    return lots


# ----------------------------------------------------------------------------- classeur abstrait
class Feuille:
    """Vue uniforme d'un onglet : titre, lignes (liste de listes, index 0 = ligne 1)."""

    def __init__(self, titre: str, lignes: list[list]):
        self.titre = titre
        self.lignes = lignes
        self.date = date_onglet(titre)
        self.famille = famille_onglet(titre)

    def index(self) -> dict[tuple, list[int]]:
        idx: dict[tuple, list[int]] = {}
        for i, ligne in enumerate(self.lignes):
            if not ligne:
                continue
            k = cle_lot(ligne[0] if len(ligne) > 0 else None)
            if k:
                idx.setdefault(k, []).append(i + 1)  # n° de ligne 1-based
        return idx


def ordre_onglets(feuilles: list[Feuille], date_lot: str) -> list[Feuille]:
    """AA du jour, autres AA (du plus proche au plus lointain), OUT du jour, autres OUT."""
    jjmm = date_lot[:4]

    def rang(f: Feuille):
        fam = 0 if f.famille == "AA" else 1
        if f.date == jjmm:
            return (fam, 0, 0)
        # distance en jours approximative (même mois d'abord)
        try:
            dist = abs(int(f.date[2:]) - int(jjmm[2:])) * 31 + abs(int(f.date[:2]) - int(jjmm[:2]))
        except (TypeError, ValueError):
            dist = 9999
        return (fam, 1, dist)

    return sorted([f for f in feuilles if f.famille], key=rang)


# ----------------------------------------------------------------------------- cœur
def planifier(lots: list[Lot], feuilles: list[Feuille]) -> list[dict]:
    """Pour chaque lot, décide la cellule à écrire. Retourne les opérations (une par lot)."""
    index = {f.titre: f.index() for f in feuilles}
    ops = []
    for lot in lots:
        op = {"fichier": lot.fichier, "montant_endosse": lot.montant, "statut": "introuvable",
              "onglet": None, "ligne": None, "montant_susi": None, "ecart": None, "rouge": False,
              "lien_ecrit": False, "candidats": []}
        if not lot.cle:
            op["statut"] = "nom illisible"
            ops.append(op)
            continue
        trouve = None
        for f in ordre_onglets(feuilles, lot.cle[0]):
            lignes = index[f.titre].get(lot.cle)
            if lignes:
                trouve = (f, lignes[0])
                if len(lignes) > 1:
                    op["candidats"] = [f"{f.titre}!{n}" for n in lignes]
                break
        if not trouve and lot.cle[3] > 0:
            # repli : même date, même catégorie, même index, nombre de chèques différent (50CH06 vs 48CH06)
            for f in ordre_onglets(feuilles, lot.cle[0]):
                for k, lignes in index[f.titre].items():
                    if k[0] == lot.cle[0] and k[1] == lot.cle[1] and k[3] == lot.cle[3] and k[2] != lot.cle[2]:
                        trouve = (f, lignes[0])
                        op["remarque"] = f"rapproché de {f.lignes[lignes[0]-1][0]!s} (nb chèques différent)"
                        break
                if trouve:
                    break
        if not trouve:
            # même date + même nb/idx, catégorie différente (FID / REC au lieu de PAIEMENTS)
            for f in ordre_onglets(feuilles, lot.cle[0]):
                for k, lignes in index[f.titre].items():
                    if k[0] == lot.cle[0] and k[2:] == lot.cle[2:] and k[1] != lot.cle[1]:
                        op["candidats"].append(f"{f.titre}!{lignes[0]} ({k[1]})")
            ops.append(op)
            continue
        f, n = trouve
        ligne = f.lignes[n - 1]
        susi = montant(ligne[COL_SUSI - 1] if len(ligne) >= COL_SUSI else None)
        # lot éclaté sur plusieurs lignes du même onglet (ex. 49 + 1 chèques) : E = somme des lignes
        autres = [m for m in index[f.titre].get(lot.cle, []) if m != n]
        if autres:
            parts = [susi] + [montant(f.lignes[m - 1][COL_SUSI - 1] if len(f.lignes[m - 1]) >= COL_SUSI else None) for m in autres]
            if all(x is not None for x in parts):
                susi = round(sum(parts), 2)
                op["remarque"] = f"E = somme de {len(parts)} lignes ({', '.join(f'{x:.2f}' for x in parts)})"
            op["lignes_lot"] = [n] + autres
        op.update(statut="ok", onglet=f.titre, ligne=n, montant_susi=susi,
                  f_precedent=montant(ligne[COL_ENDOSSE - 1] if len(ligne) >= COL_ENDOSSE else None))
        if susi is None:
            op["ecart"] = None
            op["rouge"] = False
            op["statut"] = "ok (E vide ou formule)"
        else:
            op["ecart"] = round(susi - lot.montant, 2)
            op["rouge"] = abs(op["ecart"]) >= 0.005
        lien_actuel = ligne[COL_LIEN - 1] if len(ligne) >= COL_LIEN else None
        entete = f.lignes[0][COL_LIEN - 1] if f.lignes and len(f.lignes[0]) >= COL_LIEN else None
        if lot.lien and not lien_actuel and entete and str(entete).strip().upper() in ("LIEN", "LINK"):
            op["lien_ecrit"] = True
            op["lien"] = lot.lien
        ops.append(op)
    return ops


# ----------------------------------------------------------------------------- backend xlsx
def charger_xlsx(chemin: Path):
    import openpyxl
    wb = openpyxl.load_workbook(chemin)
    feuilles = []
    for ws in wb.worksheets:
        if not famille_onglet(ws.title):
            continue
        lignes = [list(r) for r in ws.iter_rows(values_only=True)]
        feuilles.append(Feuille(ws.title, lignes))
    return wb, feuilles


def appliquer_xlsx(wb, ops: list[dict], sortie: Path):
    from openpyxl.styles import PatternFill
    rouge = PatternFill(start_color=ROUGE, end_color=ROUGE, fill_type="solid")
    for op in ops:
        if not op["onglet"]:
            continue
        ws = wb[op["onglet"]]
        n = op["ligne"]
        ws.cell(row=n, column=COL_ENDOSSE, value=op["montant_endosse"]).number_format = "#,##0.00 €"
        if op.get("lien_ecrit"):
            ws.cell(row=n, column=COL_LIEN, value=op["lien"])
        for lig in (op.get("lignes_lot") or [n]) if op["rouge"] else []:
            for c in range(1, max(COL_LIEN, ws.max_column) + 1):
                ws.cell(row=lig, column=c).fill = rouge
    wb.save(sortie)


# ----------------------------------------------------------------------------- backend Google Sheets
def ouvrir_gsheet(sheet_id: str):
    import gspread
    gc = gspread.oauth()  # ~/.config/gspread/credentials.json (client OAuth Desktop) + authorized_user.json
    return gc.open_by_key(sheet_id)


def charger_gsheet(sh):
    feuilles, wss = [], {}
    for ws in sh.worksheets():
        if not famille_onglet(ws.title):
            continue
        feuilles.append(Feuille(ws.title, ws.get_all_values()))
        wss[ws.title] = ws
    return wss, feuilles


def requetes_batch(ops: list[dict], sheet_ids: dict[str, int]) -> list[dict]:
    """Requêtes Sheets API batchUpdate (F, lien, fond rouge) ; sheet_ids = {titre onglet: sheetId}."""
    requetes = []
    for op in ops:
        if not op["onglet"]:
            continue
        if op["onglet"] not in sheet_ids:
            raise SystemExit(f"sheetId inconnu pour l'onglet {op['onglet']!r} : complète --sheet-ids")
        sid = sheet_ids[op["onglet"]]
        n = op["ligne"] - 1  # 0-based pour l'API
        requetes.append({"updateCells": {
            "range": {"sheetId": sid, "startRowIndex": n, "endRowIndex": n + 1,
                      "startColumnIndex": COL_ENDOSSE - 1, "endColumnIndex": COL_ENDOSSE},
            "rows": [{"values": [{"userEnteredValue": {"numberValue": op["montant_endosse"]},
                                  "userEnteredFormat": {"numberFormat": {"type": "NUMBER", "pattern": "#,##0.00 €"}}}]}],
            "fields": "userEnteredValue,userEnteredFormat.numberFormat"}})
        if op.get("lien_ecrit"):
            requetes.append({"updateCells": {
                "range": {"sheetId": sid, "startRowIndex": n, "endRowIndex": n + 1,
                          "startColumnIndex": COL_LIEN - 1, "endColumnIndex": COL_LIEN},
                "rows": [{"values": [{"userEnteredValue": {"stringValue": op["lien"]}}]}],
                "fields": "userEnteredValue"}})
        for lig in (op.get("lignes_lot") or [op["ligne"]]) if op["rouge"] else []:
            n = lig - 1
            requetes.append({"repeatCell": {
                "range": {"sheetId": sid, "startRowIndex": n, "endRowIndex": n + 1,
                          "startColumnIndex": 0, "endColumnIndex": COL_LIEN},
                "cell": {"userEnteredFormat": {"backgroundColor": {"red": 1, "green": 0, "blue": 0}}},
                "fields": "userEnteredFormat.backgroundColor"}})
    return requetes


def appliquer_gsheet(sh, wss, ops: list[dict]):
    requetes = requetes_batch(ops, {t: ws.id for t, ws in wss.items()})
    for i in range(0, len(requetes), 200):
        sh.batch_update({"requests": requetes[i:i + 200]})
    return len(requetes)


def ecrire_batch(ops: list[dict], sheet_ids_path: Path, dossier: Path) -> int:
    """Mode Zapier : écrit dossier/batch_NN.json = corps POST de spreadsheets/<ID>:batchUpdate (≤ 200 requêtes)."""
    sheet_ids = json.loads(sheet_ids_path.read_text(encoding="utf-8"))
    if isinstance(sheet_ids, dict) and "sheets" in sheet_ids:  # réponse brute de GET spreadsheets?fields=sheets.properties
        sheet_ids = {sh["properties"]["title"]: sh["properties"]["sheetId"] for sh in sheet_ids["sheets"]}
    requetes = requetes_batch(ops, sheet_ids)
    dossier.mkdir(parents=True, exist_ok=True)
    for i in range(0, len(requetes), 200):
        (dossier / f"batch_{i // 200 + 1:02d}.json").write_text(
            json.dumps({"requests": requetes[i:i + 200]}, ensure_ascii=False), encoding="utf-8")
    return len(requetes)


# ----------------------------------------------------------------------------- CLI
def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--rapport", required=True, type=Path, help="rapport.csv / rapport_sheet.csv de agent_ftp.py finaliser")
    src = p.add_mutually_exclusive_group(required=True)
    src.add_argument("--xlsx", type=Path, help="export xlsx du Google Sheet de suivi")
    src.add_argument("--gsheet", help="ID du Google Sheet de suivi (écriture en place via gspread/OAuth)")
    p.add_argument("--sortie", type=Path, help="xlsx patché (mode --xlsx) ; défaut : <xlsx>_rempli.xlsx")
    p.add_argument("--json", type=Path, help="écrit le détail des opérations (une par lot)")
    p.add_argument("--dry-run", action="store_true", help="ne rien écrire, seulement le plan")
    p.add_argument("--sheet-ids", type=Path, help="JSON {titre onglet: sheetId} (ou réponse GET spreadsheets?fields=sheets.properties)")
    p.add_argument("--batch-out", type=Path, help="avec --sheet-ids : dossier où écrire les corps batchUpdate (mode Zapier)")
    a = p.parse_args(argv)

    lots = lire_rapport(a.rapport)
    if a.xlsx:
        wb, feuilles = charger_xlsx(a.xlsx)
    else:
        sh = ouvrir_gsheet(a.gsheet)
        wss, feuilles = charger_gsheet(sh)
    ops = planifier(lots, feuilles)

    if not a.dry_run:
        if a.xlsx:
            sortie = a.sortie or a.xlsx.with_name(a.xlsx.stem + "_rempli.xlsx")
            appliquer_xlsx(wb, ops, sortie)
            print(f"→ {sortie}")
        else:
            n = appliquer_gsheet(sh, wss, ops)
            print(f"→ {n} requête(s) envoyées au Google Sheet {a.gsheet}")
    if a.sheet_ids and a.batch_out:
        n = ecrire_batch(ops, a.sheet_ids, a.batch_out)
        print(f"→ {n} requête(s) batchUpdate dans {a.batch_out}/batch_*.json")
    if a.json:
        a.json.write_text(json.dumps(ops, ensure_ascii=False, indent=1), encoding="utf-8")

    ok = [o for o in ops if o["onglet"]]
    rouges = [o for o in ok if o["rouge"]]
    manquants = [o for o in ops if not o["onglet"]]
    print(f"{len(ok)}/{len(ops)} lots placés — {len(rouges)} écart(s) F≠E (ligne rouge) — {len(manquants)} introuvable(s)")
    for o in ok:
        flag = f"  ⚠ écart {o['ecart']:+.2f} €" if o["rouge"] else ""
        susi = "" if o["montant_susi"] is None else f"  E={o['montant_susi']:.2f}"
        print(f"  {o['onglet']:<10} L{o['ligne']:<5} {o['fichier']:<45} F={o['montant_endosse']:.2f}{susi}{flag}"
              + ("  (+lien)" if o.get("lien_ecrit") else "") + (f"  [{o['remarque']}]" if o.get("remarque") else ""))
    for o in manquants:
        print(f"  INTROUVABLE {o['fichier']} ({o['statut']})" + (f" — proches : {', '.join(o['candidats'])}" if o["candidats"] else ""))
    return 0


if __name__ == "__main__":
    sys.exit(main())
